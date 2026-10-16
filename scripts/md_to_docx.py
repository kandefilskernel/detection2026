"""Convertit un document Markdown du dossier docs/ en Word (.docx) pour le mémoire.

Pris en charge : titres, paragraphes (gras, italique, code), listes, tableaux, citations, images, saut de page
avant chaque PARTIE, formules LaTeX $…$ et $$…$$ converties en équations Word natives (modifiables), si
Microsoft Office est installé (MML2OMML.XSL) ; sinon les formules restent en texte.

Usage : python scripts/md_to_docx.py docs/memoire_analyse_methodologie_conception.md [sortie.docx]
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

XSL_CANDIDATES = [Path(r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL"),
                  Path(r"C:\Program Files (x86)\Microsoft Office\root\Office16\MML2OMML.XSL")]
_xslt = None


def omml(latex: str):
    """Formule LaTeX -> élément OMML (équation Word), ou None si impossible."""
    global _xslt
    try:
        from latex2mathml.converter import convert
        from lxml import etree
        if _xslt is None:
            xsl = next((p for p in XSL_CANDIDATES if p.exists()), None)
            if xsl is None:
                return None
            _xslt = etree.XSLT(etree.parse(str(xsl)))
        latex = (latex.replace(r"\operatorname", r"\mathrm").replace(r"\tfrac", r"\frac").replace(r"\dfrac", r"\frac")
                 .replace(r"\!", "").replace(r"\,", " ").replace(r"\;", " ").replace("{,}", ","))
        mml = etree.fromstring(convert(latex).encode())
        return _xslt(mml).getroot()
    except Exception:  # noqa: BLE001 — formule non convertible : rendu en texte
        return None


INLINE = re.compile(r"(\$[^$]+\$|\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`)")


def add_inline(par, text: str, bold=False, italic=False, size=None):
    text = re.sub(r"<(https?://[^>]+)>", lambda m: m.group(1), text)      # liens Markdown <url> -> url
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("$") and part.endswith("$") and len(part) > 2:
            el = omml(part[1:-1])
            if el is not None:
                par._p.append(el)
                continue
            run = par.add_run(part[1:-1]); run.italic = True
        elif part.startswith("**"):
            add_inline(par, part[2:-2], bold=True, italic=italic, size=size); continue
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            add_inline(par, part[1:-1], bold=bold, italic=True, size=size); continue
        elif part.startswith("`"):
            run = par.add_run(part[1:-1]); run.font.name = "Consolas"; run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x1F, 0x3A, 0x8A)
        else:
            run = par.add_run(part)
        run.bold = run.bold or bold
        run.italic = run.italic or italic
        if size:
            run.font.size = size


def shade(cell, hex_color: str):
    tc = cell._tc.get_or_add_tcPr()
    shd = tc.makeelement(qn("w:shd"), {qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): hex_color})
    tc.append(shd)


def add_table(doc, rows: list[list[str]]):
    header, body = rows[0], rows[1:]
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ""
        add_inline(c.paragraphs[0], h, bold=True, size=Pt(9))
        shade(c, "DCE6F7")
    for r in body:
        cells = t.add_row().cells
        for i in range(len(header)):
            cells[i].text = ""
            add_inline(cells[i].paragraphs[0], r[i] if i < len(r) else "", size=Pt(9))
    doc.add_paragraph()


def split_row(line: str) -> list[str]:
    line = line.strip().strip("|")
    return [c.strip().replace("\\|", "|") for c in re.split(r"(?<!\\)\|", line)]


def convert(src: Path, out: Path):
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.2); s.top_margin = s.bottom_margin = Cm(2)
    lines = src.read_text(encoding="utf-8").splitlines()
    i, first_part = 0, True
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if not s or s == "---":
            i += 1; continue
        if s.startswith("$$"):
            buf = s[2:]
            while not buf.rstrip().endswith("$$") and i + 1 < len(lines):
                i += 1; buf += " " + lines[i].strip()
            latex = buf.rstrip()[:-2].strip()
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            el = omml(latex)
            if el is not None:
                p._p.append(el)
            else:
                r = p.add_run(latex); r.italic = True
            i += 1; continue
        m = re.match(r"^(#{1,4})\s+(.*)", s)
        if m:
            level, title = len(m.group(1)), m.group(2)
            if (level == 1 and title.startswith("PARTIE")) or (level == 2 and title.startswith("SECTION")):
                if not first_part:
                    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                first_part = False
            if level == 1 and not title.startswith("PARTIE") and i < 3:
                doc.add_heading(title, level=0)
            else:
                doc.add_heading(re.sub(r"[*`]", "", title), level=min(level, 3) if not title.startswith("PARTIE") else 1)
            i += 1; continue
        if s.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not re.match(r"^\|[\s:|-]+\|$", lines[i].strip()):
                    rows.append(split_row(lines[i]))
                i += 1
            add_table(doc, rows); continue
        img = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)", s)
        if img:
            path = (src.parent / img.group(2)).resolve()
            if path.exists():
                p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(path), width=Cm(16))
            i += 1; continue
        if s.startswith(">"):
            p = doc.add_paragraph(style="Intense Quote")
            add_inline(p, s.lstrip("> ").strip()); i += 1; continue
        lm = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)", line)
        if lm:
            style = "List Number" if lm.group(2)[0].isdigit() else "List Bullet"
            if len(lm.group(1)) >= 2:
                style += " 2"
            p = doc.add_paragraph(style=style)
            add_inline(p, lm.group(3)); i += 1; continue
        if s.startswith("*Figure") or s.startswith("*Figure"):
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(p, s.strip("*"), italic=True, size=Pt(9)); i += 1; continue
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6)
        add_inline(p, s); i += 1
    doc.save(out)
    print(f"écrit : {out}")


if __name__ == "__main__":
    src = Path(sys.argv[1])
    convert(src, Path(sys.argv[2]) if len(sys.argv) > 2 else src.with_suffix(".docx"))
